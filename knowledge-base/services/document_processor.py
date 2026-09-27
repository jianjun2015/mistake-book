"""
文档处理器 - 多格式文档解析
支持: txt, md, pdf, docx, pptx, xlsx, csv, json, xml, html, 图片OCR, 代码
"""
import os
import uuid
import json
import mimetypes
from typing import List, Dict, Any, Optional
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedResult:
    """解析结果"""
    content: str
    metadata: Dict[str, Any]


class DocumentChunker:
    """文本分块器"""
    
    def __init__(self, chunk_size: int = 500, overlap: int = 50):
        self.chunk_size = chunk_size
        self.overlap = overlap
    
    def chunk(self, text: str, metadata: Dict[str, Any] = None) -> List[Dict[str, Any]]:
        """
        智能分块策略：
        1. 优先按段落分割
        2. 段落过长按句子分割
        3. 保留overlap确保上下文连贯
        """
        if not text.strip():
            return []
        
        # 按段落分割
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        
        chunks = []
        current_chunk = ""
        chunk_index = 0
        char_pos = 0
        
        for para in paragraphs:
            # 如果段落本身就超长，按句子切分
            if len(para) > self.chunk_size:
                # 先把当前累积的入队
                if current_chunk:
                    chunks.append(self._make_chunk(current_chunk, chunk_index, metadata))
                    chunk_index += 1
                    current_chunk = ""
                
                # 长段落按句切
                sentences = self._split_sentences(para)
                for sent in sentences:
                    if len(current_chunk) + len(sent) > self.chunk_size and current_chunk:
                        chunks.append(self._make_chunk(current_chunk, chunk_index, metadata))
                        chunk_index += 1
                        # 保留overlap
                        current_chunk = current_chunk[-self.overlap:] if self.overlap else ""
                    current_chunk += sent
            else:
                if len(current_chunk) + len(para) > self.chunk_size and current_chunk:
                    chunks.append(self._make_chunk(current_chunk, chunk_index, metadata))
                    chunk_index += 1
                    current_chunk = current_chunk[-self.overlap:] if self.overlap else ""
                current_chunk += ("\n\n" if current_chunk else "") + para
        
        if current_chunk.strip():
            chunks.append(self._make_chunk(current_chunk, chunk_index, metadata))
        
        return chunks
    
    def _make_chunk(self, content: str, index: int, metadata: Dict = None) -> Dict[str, Any]:
        return {
            "id": str(uuid.uuid4()),
            "content": content.strip(),
            "chunk_index": index,
            "metadata": (metadata or {}).copy()
        }
    
    def _split_sentences(self, text: str) -> List[str]:
        import re
        parts = re.split(r'(?<=[。！？\.!\?\n])', text)
        return [p for p in parts if p.strip()]


class DocumentProcessor:
    """文档处理器 - 根据文件类型分派到对应解析器"""
    
    PARSER_MAP = {
        ".txt": "_parse_text",
        ".md": "_parse_text",
        ".log": "_parse_text",
        ".pdf": "_parse_pdf",
        ".docx": "_parse_docx",
        ".doc": "_parse_docx",
        ".pptx": "_parse_pptx",
        ".xlsx": "_parse_xlsx",
        ".xls": "_parse_xlsx",
        ".csv": "_parse_csv",
        ".json": "_parse_json",
        ".xml": "_parse_xml",
        ".html": "_parse_html",
        ".htm": "_parse_html",
        ".jpg": "_parse_image",
        ".jpeg": "_parse_image",
        ".png": "_parse_image",
        ".bmp": "_parse_image",
        ".gif": "_parse_image",
        ".webp": "_parse_image",
        ".mp3": "_parse_audio",
        ".wav": "_parse_audio",
        ".ogg": "_parse_audio",
        ".mp4": "_parse_video",
        ".avi": "_parse_video",
        ".mov": "_parse_video",
        ".py": "_parse_code",
        ".java": "_parse_code",
        ".js": "_parse_code",
        ".ts": "_parse_code",
        ".go": "_parse_code",
        ".rs": "_parse_code",
        ".cpp": "_parse_code",
        ".c": "_parse_code",
        ".sh": "_parse_code",
    }
    
    def __init__(self, chunk_size: int = 500, overlap: int = 50):
        self.chunker = DocumentChunker(chunk_size, overlap)
    
    def process_file(self, file_path: str, metadata: Dict[str, Any] = None) -> Dict[str, Any]:
        """
        处理文件，返回：
        {
            "content": str,
            "chunks": List[Dict],
            "metadata": Dict
        }
        """
        ext = Path(file_path).suffix.lower()
        parser_name = self.PARSER_MAP.get(ext, "_parse_text")
        parser = getattr(self, parser_name, self._parse_text)
        
        try:
            result = parser(file_path)
        except Exception as e:
            result = ParsedResult(content=f"[解析失败: {str(e)}]", metadata={"error": str(e)})
        
        meta = metadata or {}
        meta.setdefault("filename", os.path.basename(file_path))
        meta.setdefault("file_type", ext.lstrip("."))
        meta.setdefault("file_size", os.path.getsize(file_path))
        meta.update(result.metadata)
        
        chunks = self.chunker.chunk(result.content, meta)
        
        return {
            "content": result.content,
            "chunks": chunks,
            "metadata": meta
        }
    
    def process_text(self, text: str, metadata: Dict[str, Any] = None) -> Dict[str, Any]:
        """处理直接输入的文本"""
        meta = metadata or {}
        chunks = self.chunker.chunk(text, meta)
        return {"content": text, "chunks": chunks, "metadata": meta}
    
    # ========== 文本类 ==========
    
    def _parse_text(self, path: str) -> ParsedResult:
        import chardet
        with open(path, "rb") as f:
            raw = f.read()
        det = chardet.detect(raw)
        encoding = det.get("encoding") or "utf-8"
        text = raw.decode(encoding, errors="replace")
        return ParsedResult(content=text, metadata={"encoding": encoding})
    
    def _parse_pdf(self, path: str) -> ParsedResult:
        """PDF解析"""
        try:
            import pdfplumber
            texts = []
            page_count = 0
            with pdfplumber.open(path) as pdf:
                page_count = len(pdf.pages)
                for page in pdf.pages:
                    t = page.extract_text() or ""
                    if t.strip():
                        texts.append(t)
            return ParsedResult(
                content="\n\n".join(texts),
                metadata={"page_count": page_count}
            )
        except ImportError:
            # 回退到pypdf
            try:
                from pypdf import PdfReader
                reader = PdfReader(path)
                texts = [p.extract_text() or "" for p in reader.pages]
                return ParsedResult(
                    content="\n\n".join(texts),
                    metadata={"page_count": len(reader.pages)}
                )
            except ImportError:
                return ParsedResult(content="[PDF解析库未安装]", metadata={"error": "no pdf lib"})
    
    def _parse_docx(self, path: str) -> ParsedResult:
        ext = Path(path).suffix.lower()
        # .doc 旧格式先转为 .docx 或用 antiword
        if ext == ".doc":
            # 方案1: 用 antiword 提取文本
            import subprocess
            try:
                result = subprocess.run(["antiword", path], capture_output=True, text=True, timeout=30)
                if result.returncode == 0 and result.stdout.strip():
                    return ParsedResult(
                        content=result.stdout,
                        metadata={"format": "doc", "parser": "antiword"}
                    )
            except (FileNotFoundError, subprocess.TimeoutExpired):
                pass
            # 方案2: 用 textract
            try:
                import textract
                text = textract.process(path).decode("utf-8")
                return ParsedResult(
                    content=text,
                    metadata={"format": "doc", "parser": "textract"}
                )
            except Exception:
                pass
            # 方案3: 用 python-docx 尝试（部分 .doc 可以打开）
            try:
                from docx import Document
                doc = Document(path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                return ParsedResult(
                    content="\n\n".join(paragraphs),
                    metadata={"format": "doc", "parser": "python-docx"}
                )
            except Exception as e:
                return ParsedResult(content=f"[.doc格式解析失败: {str(e)}]", metadata={"error": str(e)})
        
        # .docx 标准格式
        try:
            from docx import Document
            doc = Document(path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # 表格
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text for c in row.cells)
                    if row_text.strip():
                        paragraphs.append(row_text)
            return ParsedResult(
                content="\n\n".join(paragraphs),
                metadata={"paragraph_count": len(paragraphs)}
            )
        except ImportError:
            return ParsedResult(content="[python-docx未安装]", metadata={"error": "no docx lib"})
    
    def _parse_pptx(self, path: str) -> ParsedResult:
        try:
            from pptx import Presentation
            prs = Presentation(path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = "".join(run.text for run in para.runs)
                            if t.strip():
                                texts.append(t)
            return ParsedResult(
                content="\n".join(texts),
                metadata={"slide_count": len(prs.slides)}
            )
        except ImportError:
            return ParsedResult(content="[python-pptx未安装]", metadata={"error": "no pptx lib"})
    
    def _parse_xlsx(self, path: str) -> ParsedResult:
        ext = Path(path).suffix.lower()
        # .xls 旧格式用 xlrd，.xlsx 用 openpyxl
        if ext == ".xls":
            try:
                import xlrd
                wb = xlrd.open_workbook(path)
                lines = []
                for sheet in wb.sheets():
                    lines.append(f"[Sheet: {sheet.name}]")
                    for row_idx in range(sheet.nrows):
                        row_text = " | ".join(str(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols) if sheet.cell_value(row_idx, col_idx))
                        if row_text.strip():
                            lines.append(row_text)
                return ParsedResult(
                    content="\n".join(lines),
                    metadata={"sheet_count": wb.nsheets, "format": "xls"}
                )
            except ImportError:
                return ParsedResult(content="[xlrd未安装，无法解析.xls文件]", metadata={"error": "no xlrd"})
        else:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(path, data_only=True)
                lines = []
                for sheet in wb.worksheets:
                    lines.append(f"[Sheet: {sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " | ".join(str(c) for c in row if c is not None)
                        if row_text.strip():
                            lines.append(row_text)
                return ParsedResult(
                    content="\n".join(lines),
                    metadata={"sheet_count": len(wb.worksheets)}
                )
            except ImportError:
                return ParsedResult(content="[openpyxl未安装]", metadata={"error": "no openpyxl"})
    
    def _parse_csv(self, path: str) -> ParsedResult:
        import csv
        import chardet
        with open(path, "rb") as f:
            raw = f.read()
        det = chardet.detect(raw)
        encoding = det.get("encoding") or "utf-8"
        lines = []
        with open(path, "r", encoding=encoding, errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                lines.append(" | ".join(row))
        return ParsedResult(content="\n".join(lines), metadata={"encoding": encoding})
    
    def _parse_json(self, path: str) -> ParsedResult:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        # 转为格式化文本
        text = json.dumps(data, ensure_ascii=False, indent=2)
        return ParsedResult(content=text, metadata={"json_type": type(data).__name__})
    
    def _parse_xml(self, path: str) -> ParsedResult:
        try:
            from bs4 import BeautifulSoup
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                soup = BeautifulSoup(f, "xml")
            return ParsedResult(content=soup.get_text(separator="\n"), metadata={})
        except ImportError:
            return self._parse_text(path)
    
    def _parse_html(self, path: str) -> ParsedResult:
        try:
            from bs4 import BeautifulSoup
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                soup = BeautifulSoup(f, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header"]):
                tag.decompose()
            title = soup.title.string if soup.title else ""
            return ParsedResult(
                content=soup.get_text(separator="\n"),
                metadata={"title": title}
            )
        except ImportError:
            return self._parse_text(path)
    
    # ========== 图片（OCR） ==========
    
    def _parse_image(self, path: str) -> ParsedResult:
        """图片OCR识别（可选依赖，按需安装）"""
        try:
            from PIL import Image
            # 先尝试pytesseract
            try:
                import pytesseract
                img = Image.open(path)
                text = pytesseract.image_to_string(img, lang="chi_sim+eng")
                return ParsedResult(
                    content=text,
                    metadata={"ocr_engine": "tesseract", "size": img.size}
                )
            except ImportError:
                # 回退：仅返回图片元信息
                img = Image.open(path)
                return ParsedResult(
                    content=f"[图片: {os.path.basename(path)}, 尺寸: {img.size}, 模式: {img.mode}]",
                    metadata={"ocr_engine": "none", "size": img.size}
                )
        except ImportError:
            return ParsedResult(content=f"[图片: {os.path.basename(path)}]", metadata={"ocr_engine": "none"})
    
    # ========== 音视频（语音转文字） ==========
    
    def _parse_audio(self, path: str) -> ParsedResult:
        """音频转文字（可选）"""
        try:
            import whisper
            model = whisper.load_model("base")
            result = model.transcribe(path, language="zh")
            return ParsedResult(
                content=result["text"],
                metadata={"asr_engine": "whisper", "language": result.get("language", "zh")}
            )
        except ImportError:
            return ParsedResult(content=f"[音频: {os.path.basename(path)}]", metadata={"asr_engine": "none"})
    
    def _parse_video(self, path: str) -> ParsedResult:
        """视频提取音频转文字（可选）"""
        return self._parse_audio(path)
    
    # ========== 代码 ==========
    
    def _parse_code(self, path: str) -> ParsedResult:
        result = self._parse_text(path)
        ext = Path(path).suffix.lower().lstrip(".")
        result.metadata["language"] = ext
        # 添加代码说明
        content = f"[代码文件: {os.path.basename(path)}, 语言: {ext}]\n{result.content}"
        return ParsedResult(content=content, metadata=result.metadata)
